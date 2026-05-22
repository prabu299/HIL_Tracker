"""
HIL Rig Tracker — Streamlit Web App
====================================
pip install streamlit matplotlib python-pptx supabase
streamlit run streamlit_app.py
"""

import streamlit as st
import streamlit.components.v1 as components
import json, datetime, os, io, tempfile
from copy import deepcopy
from collections import defaultdict

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.ticker as ticker
import numpy as np

from pptx import Presentation as PptxPrs
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ══════════════════════════════════════════════════════════════════════════════
# CONSTANTS  (identical to desktop app)
# ══════════════════════════════════════════════════════════════════════════════
DAY_HOURS  = 24
HUMAN_POOL = 18
AUTDEV_MAX = 18

CATS = [
    {"key":"autoExec","label":"Auto Execution", "short":"AUTO","hex":"#4a90d9","rgb":(74,144,217)},
    {"key":"manExec", "label":"Manual Execution","short":"MAN", "hex":"#3fb950","rgb":(63,185,80)},
    {"key":"autoDev", "label":"Auto Development","short":"DEV", "hex":"#bc8cff","rgb":(188,140,255)},
    {"key":"idle",    "label":"Idle",            "short":"IDLE","hex":"#d29922","rgb":(210,153,34)},
    {"key":"down",    "label":"Down / Inactive", "short":"DOWN","hex":"#888899","rgb":(136,136,153)},
]
CAT_KEYS   = [c["key"] for c in CATS]
CAT_BY_KEY = {c["key"]: c for c in CATS}

IDLE_PRESETS = ["Waiting for test case","No test scheduled","Resource unavailable",
                "Environment setup","Team meeting / stand-up","Other"]
DOWN_PRESETS = ["Hardware failure","NIC / network issue","Power failure","Software crash",
                "Maintenance / upgrade","Awaiting spare part","GUT not installed",
                "Dashboard issue","Other"]

STATUS_OPTS = ["active","inactive","in_progress"]
STATUS_COLS = {"active":"#3fb950","inactive":"#888899","in_progress":"#d29922"}

DEFAULT_RIGS = [{"id":f"hil-slot-{i+1}","name":f"HIL-{i+1:02d}","note":"","status":"active"}
                for i in range(7)]

# ── Chart colours
C_AUTO = "#388bfd";  C_MAN = "#3fb950";  C_DEV = "#bc8cff"
C_IDLE = "#d29922";  C_DOWN = "#888899"; C_ACCENT = "#388bfd"
C_GREEN = "#3fb950"; C_RED = "#f85149";  C_YELLOW = "#d29922"
CAT_COLORS = [C_AUTO, C_MAN, C_DEV, C_IDLE, C_DOWN]

# ── Matplotlib dark style
plt.rcParams.update({
    "font.family":"DejaVu Sans","font.size":11,
    "text.color":"#ffffff","axes.labelcolor":"#e6edf3",
    "xtick.color":"#e6edf3","ytick.color":"#e6edf3",
    "xtick.labelsize":11,"ytick.labelsize":11,
    "axes.facecolor":"#1c2128","figure.facecolor":"#0d1117",
    "axes.edgecolor":"#2d3748","grid.color":"#2d3748",
    "legend.facecolor":"#21262d","legend.edgecolor":"#30363d",
    "legend.labelcolor":"#e6edf3","legend.fontsize":11,
})

# ══════════════════════════════════════════════════════════════════════════════
# STREAMLIT PAGE CONFIG
# ══════════════════════════════════════════════════════════════════════════════
st.set_page_config(
    page_title="HIL Rig Tracker",
    page_icon="🔬",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ── Inject CSS to match dark desktop theme ───────────────────────────────────
st.markdown("""
<style>
/* Dark background everywhere */
.stApp, .main, section[data-testid="stSidebar"] {
    background-color: #0d1117 !important;
    color: #e6edf3 !important;
}
/* Header */
header[data-testid="stHeader"] { background: #161b22 !important; border-bottom: 1px solid #30363d; }
/* Sidebar */
section[data-testid="stSidebar"] > div:first-child { background: #161b22 !important; border-right: 1px solid #30363d; }
/* Tabs */
.stTabs [data-baseweb="tab-list"] { background: #161b22; border-bottom: 1px solid #30363d; gap: 0; }
.stTabs [data-baseweb="tab"] {
    background: transparent; color: #8b949e; border-bottom: 2px solid transparent;
    padding: 10px 24px; font-size: 11pt; font-weight: 500;
}
.stTabs [aria-selected="true"] { color: #e6edf3 !important; border-bottom: 2px solid #388bfd !important; }
.stTabs [data-baseweb="tab-panel"] { background: #0d1117; padding-top: 16px; }
/* Buttons */
.stButton > button {
    background: #21262d; color: #e6edf3; border: 1px solid #30363d;
    border-radius: 6px; font-size: 10pt;
}
.stButton > button:hover { background: #30363d; border-color: #8b949e; }
/* Inputs */
.stTextInput input, .stNumberInput input, .stDateInput input,
.stSelectbox > div > div, .stTextArea textarea {
    background: #161b22 !important; color: #e6edf3 !important;
    border: 1px solid #30363d !important; border-radius: 6px !important;
}
/* Number inputs */
input[type=number] { background: #161b22 !important; color: #e6edf3 !important; }
/* Dataframes / tables */
.dataframe, [data-testid="stDataFrame"] {
    background: #161b22 !important; color: #e6edf3 !important;
}
/* Metric cards */
[data-testid="stMetricValue"] { color: #e6edf3; font-size: 24pt; font-weight: 700; }
[data-testid="stMetricLabel"] { color: #8b949e; font-size: 9pt; letter-spacing: 0.8px; }
[data-testid="metric-container"] {
    background: #161b22; border: 1px solid #30363d; border-radius: 8px; padding: 12px 16px;
}
/* Expander */
.streamlit-expanderHeader { background: #161b22 !important; color: #e6edf3 !important; border: 1px solid #30363d; }
.streamlit-expanderContent { background: #1c2128 !important; border: 1px solid #30363d; border-top: none; }
/* Checkbox / radio */
.stCheckbox label, .stRadio label { color: #e6edf3 !important; }
/* Alert/info boxes */
.stAlert { background: #1c2128; border: 1px solid #30363d; }
/* Selectbox dropdown */
[data-baseweb="select"] > div { background: #161b22 !important; border-color: #30363d !important; }
/* Hide hamburger and footer */
#MainMenu { visibility: hidden; }
footer { visibility: hidden; }
/* Custom card style */
.rig-card {
    background: #161b22; border: 1px solid #30363d; border-radius: 8px;
    padding: 12px 16px; margin: 4px 0;
}
/* Success/warning inline */
.tag-active   { color:#3fb950; background:rgba(63,185,80,0.12); border:1px solid rgba(63,185,80,0.3); border-radius:4px; padding:2px 8px; font-size:9pt; font-weight:700; }
.tag-inactive { color:#888899; background:rgba(136,136,153,0.12); border:1px solid rgba(136,136,153,0.3); border-radius:4px; padding:2px 8px; font-size:9pt; font-weight:700; }
.tag-progress { color:#d29922; background:rgba(210,153,34,0.12); border:1px solid rgba(210,153,34,0.3); border-radius:4px; padding:2px 8px; font-size:9pt; font-weight:700; }
</style>
""", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# DATABASE LAYER  (Supabase → falls back to session_state JSON)
# ══════════════════════════════════════════════════════════════════════════════
@st.cache_resource
def get_supabase():
    """Return Supabase client or None if not configured."""
    try:
        url = st.secrets["SUPABASE_URL"]
        key = st.secrets["SUPABASE_KEY"]
        from supabase import create_client
        return create_client(url, key)
    except Exception:
        return None

def db_load():
    """Load all data from Supabase. Returns {logs:{}, rigs:[]}."""
    sb = get_supabase()
    if sb is None:
        return None
    try:
        rows = sb.table("hil_tracker").select("*").execute().data or []
        result = {"logs": {}, "rigs": None}
        for row in rows:
            try:
                data = json.loads(row["log_data"])
                if row["record_type"] == "rigs":
                    result["rigs"] = data
                elif row["record_type"] == "log":
                    date = row["title"].replace("log__", "")
                    result["logs"][date] = data
            except Exception:
                pass
        return result
    except Exception:
        return None

def db_upsert(title, record_type, data_obj):
    """Insert or update a single record."""
    sb = get_supabase()
    if sb is None:
        return
    try:
        sb.table("hil_tracker").upsert({
            "title": title,
            "record_type": record_type,
            "log_data": json.dumps(data_obj),
        }, on_conflict="title").execute()
    except Exception as e:
        st.warning(f"DB save warning: {e}")

def db_save_log(date, day_data):
    db_upsert(f"log__{date}", "log", day_data)

def db_save_rigs(rigs):
    db_upsert("rigs__config", "rigs", rigs)

# ── Session state initialisation ─────────────────────────────────────────────
def init_state():
    if "store" not in st.session_state:
        loaded = db_load()
        if loaded and (loaded["rigs"] or loaded["logs"]):
            st.session_state.store = {
                "logs": loaded["logs"] or {},
                "rigs": loaded["rigs"] or deepcopy(DEFAULT_RIGS),
            }
        else:
            st.session_state.store = {"logs": {}, "rigs": deepcopy(DEFAULT_RIGS)}
        for r in st.session_state.store["rigs"]:
            r.setdefault("note", "")
            r.setdefault("status", "active")
    if "user_tz" not in st.session_state:
        st.session_state.user_tz = None

# ══════════════════════════════════════════════════════════════════════════════
# HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def blank_entry():
    d = {k:0.0 for k in CAT_KEYS}
    d.update({"idleReason":"","downReason":"","project":"","swVersion":"","notes":""})
    return d


def get_browser_timezone():
    if st.session_state.get("user_tz"):
        return st.session_state.user_tz
    tz = "UTC"
    try:
        tz_val = components.html(
            """
            <script src="https://cdn.jsdelivr.net/npm/streamlit-component-lib@0.0.0/dist/streamlit-component-lib.min.js"></script>
            <script>
              const tz = Intl.DateTimeFormat().resolvedOptions().timeZone || 'UTC';
              Streamlit.setComponentValue(tz);
            </script>
            """,
            height=0,
        )
        if isinstance(tz_val, str) and tz_val:
            tz = tz_val
    except Exception:
        pass
    st.session_state.user_tz = tz
    return tz


def get_user_now():
    tz_name = get_browser_timezone()
    try:
        from zoneinfo import ZoneInfo
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = datetime.timezone.utc
    return datetime.datetime.now(tz)


def today_date():
    return get_user_now().date()

def today_str():   return today_date().isoformat()
def fmt_date(s):
    try:    return datetime.date.fromisoformat(s).strftime("%d %b %Y")
    except: return s
def fmt_short(s):
    try:    return datetime.date.fromisoformat(s).strftime("%d %b")
    except: return s
def last_n_dates(n):
    t=datetime.date.today()
    return [(t-datetime.timedelta(days=i)).isoformat() for i in range(n-1,-1,-1)]
def date_range_list(f,t):
    out,cur=[],datetime.date.fromisoformat(f)
    end=datetime.date.fromisoformat(t)
    while cur<=end: out.append(cur.isoformat()); cur+=datetime.timedelta(days=1)
    return out
def sum_hours(e):
    return sum(float(e.get(k,0) or 0) for k in CAT_KEYS)

# ══════════════════════════════════════════════════════════════════════════════
# CHART HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def style_ax(ax, ylabel="", title=""):
    ax.set_facecolor("#1c2128")
    ax.tick_params(colors="#e6edf3", labelsize=11, length=0, pad=5)
    for sp in ax.spines.values(): sp.set_edgecolor("#2d3748"); sp.set_linewidth(0.6)
    ax.grid(axis="y", color="#2d3748", linewidth=0.7, alpha=0.8)
    ax.grid(axis="x", visible=False)
    if ylabel: ax.set_ylabel(ylabel, color="#e6edf3", fontsize=11, labelpad=6)
    if title:  ax.set_title(title, color="#ffffff", fontsize=13, pad=10, fontweight="bold")

def make_fig(h=4.5):
    fig, axes = plt.subplots(figsize=(10, h), facecolor="#0d1117")
    return fig, axes

def weekly_util_chart(store, sel_rigs=None):
    """Grafana-style stacked bar chart."""
    rigs  = [r for r in store["rigs"] if sel_rigs is None or r["id"] in sel_rigs]
    logs  = store["logs"]
    dates = last_n_dates(7)
    names = [r["name"] for r in rigs]
    x     = np.arange(len(rigs))

    vals = {k: [] for k in CAT_KEYS}
    for rig in rigs:
        for k in CAT_KEYS:
            vals[k].append(sum(float(logs.get(d,{}).get(rig["id"],{}).get(k,0) or 0)
                               for d in dates))

    fig, ax = plt.subplots(figsize=(10, 4.5), facecolor="#0d1117")
    style_ax(ax, ylabel="Active time (hrs)", title="Weekly Utilisation (last 7 days)")

    order  = ["autoExec","manExec","autoDev","idle","down"]
    colors = [C_AUTO, C_MAN, C_DEV, C_IDLE, C_DOWN]
    labels = ["Automated","Manual","Dev","Idle","Inactive"]
    btm    = np.zeros(len(rigs)); patches = []
    for key, col, lbl in zip(order, colors, labels):
        v = np.array(vals[key])
        ax.bar(x, v, 0.55, bottom=btm, color=col, alpha=0.9, zorder=3)
        btm += v
        patches.append(mpatches.Patch(color=col, label=lbl))

    ymax = max(50, int(np.ceil(float(np.max(btm))/50)*50)+20) if len(btm) else 50
    ax.set_ylim(0, ymax)
    ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v,_: f"{int(v)} hrs"))
    ax.set_xticks(x)
    ax.set_xticklabels(names, rotation=20, ha="right", fontsize=11, color="#c9d1d9")
    ax.tick_params(axis="y", labelsize=11, colors="#c9d1d9", pad=5)
    for xi, tot in enumerate(btm):
        if tot > 0:
            ax.text(xi, tot+1, f"{tot:.0f}h", ha="center", va="bottom",
                    fontsize=10, color="#ffffff", fontweight="bold")
    ax.legend(handles=patches, loc="upper center", bbox_to_anchor=(0.5,-0.22),
              ncol=5, frameon=True, fontsize=11, framealpha=0.9,
              facecolor="#21262d", edgecolor="#30363d", labelcolor="#e6edf3")
    fig.tight_layout(rect=[0, 0.05, 1, 1])
    return fig

def report_charts(stats, cdata):
    """Returns (fig1, fig2) for report page."""
    nms    = [r["name"] for r in stats]
    x      = np.arange(len(stats))
    dates  = [r["date"] for r in cdata]
    order  = ["autoExec","manExec","autoDev","idle","down"]
    colors = [C_AUTO, C_MAN, C_DEV, C_IDLE, C_DOWN]
    labels = ["Automated","Manual","Dev","Idle","Inactive"]

    # Fig1: weekly stacked + utilisation %
    fig1, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5.2), facecolor="#0d1117")
    fig1.subplots_adjust(wspace=0.3, left=0.06, right=0.97, top=0.92, bottom=0.18)
    style_ax(ax1, ylabel="Active time", title="Weekly Utilisation (last period)")
    btm = np.zeros(len(stats)); patches = []
    for key, col, lbl in zip(order, colors, labels):
        v = np.array([r[key] for r in stats])
        ax1.bar(x, v, 0.55, bottom=btm, color=col, alpha=0.9, zorder=3)
        btm += v
        patches.append(mpatches.Patch(color=col, label=lbl))
    ax1.set_xticks(x); ax1.set_xticklabels(nms, rotation=20, ha="right", fontsize=11, color="#c9d1d9")
    ax1.tick_params(axis="y", labelsize=11, colors="#c9d1d9")
    ymax = max(50, int(np.ceil(float(np.max(btm))/50)*50)+20) if len(btm) else 50
    ax1.set_ylim(0, ymax)
    ax1.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v,_: f"{int(v)} hrs"))
    for xi, tot in enumerate(btm):
        if tot>0: ax1.text(xi, tot+1, f"{tot:.0f}h", ha="center", va="bottom", fontsize=10, color="#fff", fontweight="bold")
    ax1.legend(handles=patches, loc="upper center", bbox_to_anchor=(0.5,-0.28),
               ncol=3, frameon=True, fontsize=10, facecolor="#21262d",
               edgecolor="#30363d", labelcolor="#e6edf3")

    style_ax(ax2, ylabel="%", title="Utilisation % per Rig")
    utils  = [r["util"] for r in stats]
    bc     = [C_GREEN if u>=70 else C_YELLOW if u>=40 else C_RED for u in utils]
    bars   = ax2.bar(nms, utils, 0.55, color=bc, alpha=0.9, zorder=3)
    ax2.set_ylim(0, 115)
    ax2.axhline(70, color=C_GREEN+"66", lw=0.8, ls="--", zorder=2)
    ax2.axhline(40, color=C_YELLOW+"66", lw=0.8, ls="--", zorder=2)
    plt.setp(ax2.get_xticklabels(), rotation=20, ha="right", fontsize=11, color="#c9d1d9")
    ax2.tick_params(axis="y", labelsize=11, colors="#c9d1d9")
    for bar, val in zip(bars, utils):
        ax2.text(bar.get_x()+bar.get_width()/2, bar.get_height()+2,
                 f"{val}%", ha="center", va="bottom", fontsize=11, color="#fff", fontweight="bold")

    # Fig2: daily trend + pie
    fig2, (ax3, ax4) = plt.subplots(1, 2, figsize=(12, 4.8), facecolor="#0d1117")
    fig2.subplots_adjust(wspace=0.3, left=0.06, right=0.97, top=0.92, bottom=0.18)
    style_ax(ax3, ylabel="Hours", title="Daily Fleet Hours by Category")
    btm2 = np.zeros(len(cdata))
    for key, col, lbl in zip(order, colors, labels):
        v = np.array([r[key] for r in cdata])
        ax3.bar(dates, v, 0.6, bottom=btm2, color=col, alpha=0.88, label=lbl, zorder=3)
        btm2 += v
    ax3.legend(fontsize=11, facecolor="#21262d", labelcolor="#e6edf3",
               loc="upper right", framealpha=0.9, edgecolor="#30363d")
    plt.setp(ax3.get_xticklabels(), rotation=40, ha="right", fontsize=11, color="#c9d1d9")
    ax3.tick_params(axis="y", labelsize=11, colors="#c9d1d9")

    ax4.set_facecolor("#1c2128")
    fleet = {k: sum(r[k] for r in stats) for k in CAT_KEYS}
    pv = [fleet[k] for k in order]; pc = colors; pl = labels
    if sum(pv) > 0:
        wedges, texts, pcts = ax4.pie(pv, labels=pl, colors=pc, autopct="%1.0f%%",
                                       startangle=90, pctdistance=0.78,
                                       textprops={"color":"#c9d1d9","fontsize":11})
        for p in pcts: p.set_color("#ffffff"); p.set_fontweight("bold")
        for t in texts: t.set_color("#c9d1d9")
    ax4.set_title("Fleet Time Distribution", color="#ffffff", fontsize=13, pad=10, fontweight="bold")

    return fig1, fig2

# ══════════════════════════════════════════════════════════════════════════════
# PPTX EXPORT  (same as desktop)
# ══════════════════════════════════════════════════════════════════════════════
def build_pptx(store, stats, cdata, dates, from_d, to_d, fig1, fig2):
    prs = PptxPrs()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def rgb(h): return RGBColor(int(h[1:3],16),int(h[3:5],16),int(h[5:7],16))
    BG_R=rgb("#0d1117"); BG2_R=rgb("#161b22"); BG3_R=rgb("#1c2128")
    FG_R=rgb("#e6edf3"); FG3_R=rgb("#8b949e"); ACC_R=rgb("#388bfd")
    GRN_R=rgb("#3fb950"); YLW_R=rgb("#d29922"); RED_R=rgb("#f85149")
    INA_R=rgb("#888899"); MAN_R=rgb("#3fb950"); DEV_R=rgb("#bc8cff")

    def set_bg(sl):
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = BG_R

    def txt(sl, text, l, t, w, h, size=10, bold=False, color=None,
            align=PP_ALIGN.LEFT, font="Segoe UI"):
        tb = sl.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
        tf = tb.text_frame; tf.word_wrap = False
        p  = tf.paragraphs[0]; p.alignment = align
        run = p.add_run(); run.text = text
        run.font.size=Pt(size); run.font.bold=bold
        run.font.name=font; run.font.color.rgb = color or FG_R

    def rect(sl, l, t, w, h, fill, border=None):
        shp = sl.shapes.add_shape(1,Inches(l),Inches(t),Inches(w),Inches(h))
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
        if border: shp.line.color.rgb=border; shp.line.width=Pt(0.5)
        else: shp.line.fill.background()

    rigs = store["rigs"]; logs = store["logs"]
    active_c   = sum(1 for r in rigs if r.get("status","active")=="active")
    inactive_c = sum(1 for r in rigs if r.get("status","active")=="inactive")
    inprog_c   = sum(1 for r in rigs if r.get("status","active")=="in_progress")

    # Slide 1: Title
    sl1 = prs.slides.add_slide(blank); set_bg(sl1)
    rect(sl1,0,0,13.33,0.5,ACC_R)
    txt(sl1,"CONTINUOUS DELIVERY FOR QUALITY AND COMPLIANCE",0.4,0.1,12,0.3,
        size=9,color=RGBColor(255,255,255),bold=True)
    rect(sl1,0.3,0.7,12.7,6.5,BG2_R,rgb("#30363d"))
    txt(sl1,"CDQC ADAS REGRESSION TESTING",0.6,1.1,12,0.9,size=24,bold=True,color=FG_R)
    txt(sl1,f"Report Period:  {fmt_date(from_d)}  →  {fmt_date(to_d)}",
        0.6,2.1,12,0.5,size=13,color=ACC_R)
    txt(sl1,f"Generated: {fmt_date(today_str())}   |   {len(rigs)} Rigs",
        0.6,2.7,12,0.4,size=10,color=FG3_R)
    rect(sl1,0.6,3.3,12.1,1.2,ACC_R)
    txt(sl1,"Active  /  Inactive  /  In Progress",0.6,3.38,12.1,0.45,
        size=12,color=RGBColor(255,255,255),align=PP_ALIGN.CENTER,bold=True)
    txt(sl1,f"{active_c}  /  {inactive_c}  /  {inprog_c}",0.6,3.82,12.1,0.6,
        size=22,bold=True,color=RGBColor(255,255,255),align=PP_ALIGN.CENTER)
    notes_rigs = [r for r in rigs if r.get("note","")]
    if notes_rigs:
        txt(sl1,"Rig Notes:",0.6,4.8,12,0.35,size=10,bold=True,color=FG3_R)
        for ni,rig in enumerate(notes_rigs[:4]):
            txt(sl1,f"•  {rig['name']} :  {rig['note']}",
                0.7,5.15+ni*0.42,12,0.38,size=10,color=rgb("#c9d1d9"))

    # Slide 2: KPIs + table
    sl2 = prs.slides.add_slide(blank); set_bg(sl2)
    rect(sl2,0,0,13.33,0.5,ACC_R)
    txt(sl2,"FLEET SUMMARY",0.4,0.1,12,0.3,size=9,color=RGBColor(255,255,255),bold=True)
    txt(sl2,f"{fmt_date(from_d)} → {fmt_date(to_d)}",0.4,0.6,12,0.4,size=13,bold=True,color=FG_R)
    fu = round(sum(r["util"] for r in stats)/len(stats)) if stats else 0
    uc_r = GRN_R if fu>=70 else YLW_R if fu>=40 else RED_R
    kpis = [("FLEET UTIL",f"{fu}%",uc_r),
            ("AUTO EXEC",f"{sum(r['autoExec'] for r in stats):.1f}h",ACC_R),
            ("MANUAL EXEC",f"{sum(r['manExec'] for r in stats):.1f}h",MAN_R),
            ("AUTO DEV",f"{sum(r['autoDev'] for r in stats):.1f}h",DEV_R),
            ("IDLE",f"{sum(r['idle'] for r in stats):.1f}h",YLW_R),
            ("INACTIVE",f"{sum(r['down'] for r in stats):.1f}h",INA_R)]
    cw=2.0; gap=0.1
    for i,(lbl,val,col) in enumerate(kpis):
        x=0.4+i*(cw+gap)
        rect(sl2,x,1.15,cw,1.6,BG3_R,rgb("#30363d"))
        txt(sl2,lbl,x+0.12,1.23,cw-0.2,0.32,size=7.5,color=FG3_R,bold=True)
        txt(sl2,val,x+0.12,1.56,cw-0.2,0.7,size=20,bold=True,color=col)
    txt(sl2,"UTILISATION BY RIG",0.4,3.0,12,0.3,size=8,color=FG3_R,bold=True)
    ch=["RIG","DAYS","AUTO EXEC","MANUAL EXEC","AUTO DEV","IDLE","DOWN","UTIL%"]
    cws=[1.5,0.65,1.2,1.3,1.2,1.0,1.0,0.8]; y_h=3.35
    for ci,(h,cw2) in enumerate(zip(ch,cws)):
        xx=0.4+sum(cws[:ci]); rect(sl2,xx,y_h,cw2-0.03,0.33,BG2_R)
        txt(sl2,h,xx+0.05,y_h+0.06,cw2-0.1,0.24,size=7.5,bold=True,color=FG3_R)
    for ri,r in enumerate(stats):
        yr=y_h+0.33+ri*0.37
        if yr>7.1: break
        rb=BG3_R if ri%2==0 else BG2_R
        uc2=GRN_R if r["util"]>=70 else YLW_R if r["util"]>=40 else RED_R
        rv=[r["name"],str(r["days"]),f"{r['autoExec']:.1f}",f"{r['manExec']:.1f}",
            f"{r['autoDev']:.1f}",f"{r['idle']:.1f}",f"{r['down']:.1f}",f"{r['util']}%"]
        rc=[FG_R,FG3_R,ACC_R,MAN_R,DEV_R,YLW_R,INA_R,uc2]
        for ci,(v,c,cw2) in enumerate(zip(rv,rc,cws)):
            xx=0.4+sum(cws[:ci]); rect(sl2,xx,yr,cw2-0.03,0.35,rb)
            txt(sl2,v,xx+0.07,yr+0.08,cw2-0.12,0.24,size=8.5,color=c)

    # Slide 3: Charts
    sl3 = prs.slides.add_slide(blank); set_bg(sl3)
    rect(sl3,0,0,13.33,0.5,ACC_R)
    txt(sl3,"USAGE CHARTS",0.4,0.1,12,0.3,size=9,color=RGBColor(255,255,255),bold=True)
    txt(sl3,f"{fmt_date(from_d)} → {fmt_date(to_d)}",0.4,0.6,8,0.4,size=13,bold=True,color=FG_R)
    def fig2png(fig):
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="#0d1117")
        buf.seek(0); return buf
    try:
        p1=fig2png(fig1); p2=fig2png(fig2)
        sl3.shapes.add_picture(p1,Inches(0.3),Inches(1.1),Inches(12.7),Inches(3.3))
        sl3.shapes.add_picture(p2,Inches(0.3),Inches(4.5),Inches(12.7),Inches(2.8))
    except Exception as e:
        txt(sl3,f"Chart error: {e}",0.5,3,12,1,size=11,color=RED_R)

    # Slide 4: Key Highlights
    sl4 = prs.slides.add_slide(blank); set_bg(sl4)
    rect(sl4,0,0,13.33,0.5,ACC_R)
    txt(sl4,"KEY HIGHLIGHTS",0.4,0.1,12,0.3,size=9,color=RGBColor(255,255,255),bold=True)
    idle_g = defaultdict(list); down_g = defaultdict(list)
    for d in dates:
        for rig in rigs:
            e=logs.get(d,{}).get(rig["id"],{})
            if not e: continue
            ih=float(e.get("idle",0)or 0); dh=float(e.get("down",0)or 0)
            ir=e.get("idleReason",""); dr=e.get("downReason","")
            if ih>0 and ir: idle_g[ir].append({"rig":rig["name"],"date":fmt_short(d),"h":ih})
            if dh>0 and dr: down_g[dr].append({"rig":rig["name"],"date":fmt_short(d),"h":dh})
    y4=1.15
    for groups,col_r,icon,sec in [(down_g,RED_R,"⚠","DOWN"),(idle_g,YLW_R,"⏸","IDLE")]:
        if not groups: continue
        sg=sorted(groups.items(),key=lambda x:sum(i["h"] for i in x[1]),reverse=True)
        th=sum(sum(i["h"] for i in v) for _,v in sg)
        txt(sl4,f"{icon} {sec} — {th:.1f}h",0.4,y4,12,0.38,size=11,bold=True,color=col_r)
        y4+=0.44
        for reason,entries in sg[:4]:
            if y4>6.9: break
            rect(sl4,0.4,y4,12.5,0.65,BG3_R,rgb("#30363d"))
            txt(sl4,f"{icon} {reason}",0.55,y4+0.05,7.5,0.28,size=10,bold=True,color=col_r)
            aff=list(set(i["rig"] for i in entries))
            txt(sl4,f"Rigs: {', '.join(aff)}",0.55,y4+0.33,12,0.26,size=8.5,color=FG3_R)
            y4+=0.71

    # Slide 5: Daily detail
    sl5=prs.slides.add_slide(blank); set_bg(sl5)
    rect(sl5,0,0,13.33,0.5,ACC_R)
    txt(sl5,"DAILY DETAIL",0.4,0.1,12,0.3,size=9,color=RGBColor(255,255,255),bold=True)
    max_days=14; chunk=dates[:max_days]
    rh=0.35; cw3=min(12.3/(len(chunk)+1),1.1); nw=1.4; y5=0.7
    for ci,d in enumerate(chunk):
        xx=nw+0.4+ci*cw3; rect(sl5,xx,y5,cw3-0.03,rh,BG2_R)
        txt(sl5,fmt_short(d),xx+0.04,y5+0.08,cw3-0.08,0.22,size=7.5,bold=True,color=FG3_R,align=PP_ALIGN.CENTER)
    for ri,rig in enumerate(rigs):
        yr=y5+rh+ri*rh
        if yr>7.1: break
        rb=BG3_R if ri%2==0 else BG2_R
        rect(sl5,0.4,yr,nw-0.03,rh,rb)
        txt(sl5,rig["name"],0.45,yr+0.09,nw-0.08,0.22,size=9,bold=True,color=FG_R)
        for ci,d in enumerate(chunk):
            xx=nw+0.4+ci*cw3; e=logs.get(d,{}).get(rig["id"],{})
            rect(sl5,xx,yr,cw3-0.03,rh,rb)
            if e:
                ah=sum(float(e.get(k,0)or 0) for k in ["autoExec","manExec","autoDev"])
                dh=float(e.get("down",0)or 0); ih=float(e.get("idle",0)or 0)
                if ah>0: tv,tc=f"{ah:.1f}h",GRN_R
                elif dh>0: tv,tc="DOWN",INA_R
                elif ih>0: tv,tc="IDLE",YLW_R
                else: tv,tc="—",FG3_R
            else: tv,tc="—",FG3_R
            txt(sl5,tv,xx+0.04,yr+0.09,cw3-0.08,0.22,size=8,color=tc,align=PP_ALIGN.CENTER)

    buf = io.BytesIO()
    prs.save(buf); buf.seek(0)
    return buf

# ══════════════════════════════════════════════════════════════════════════════
# COMPUTE REPORT STATS
# ══════════════════════════════════════════════════════════════════════════════
def compute_stats(store, from_d, to_d, sel_rig_ids=None):
    dates = date_range_list(from_d, to_d)
    logs  = store["logs"]
    rigs  = [r for r in store["rigs"] if sel_rig_ids is None or r["id"] in sel_rig_ids]
    stats = []
    for rig in rigs:
        agg={k:0.0 for k in CAT_KEYS}; days=0
        for d in dates:
            e=logs.get(d,{}).get(rig["id"],{})
            if e and any(float(e.get(k,0)or 0)>0 for k in CAT_KEYS):
                days+=1
                for k in CAT_KEYS: agg[k]+=float(e.get(k,0)or 0)
        cap=len(dates)*DAY_HOURS
        active=agg["autoExec"]+agg["manExec"]+agg["autoDev"]
        util=round(active/cap*100) if cap else 0
        stats.append({"id":rig["id"],"name":rig["name"],"days":days,"util":util,**agg})
    cdata=[]
    for d in dates:
        row={"date":fmt_short(d)}
        for k in CAT_KEYS:
            row[k]=sum(float(logs.get(d,{}).get(r["id"],{}).get(k,0)or 0) for r in rigs)
        cdata.append(row)
    return stats, cdata, dates

# ══════════════════════════════════════════════════════════════════════════════
# SIDEBAR
# ══════════════════════════════════════════════════════════════════════════════
def render_sidebar(store):
    with st.sidebar:
        st.markdown("## 🔬 HIL Rig Tracker")
        st.markdown(f"**{len(store['rigs'])} rigs** · **{len(store['logs'])} days logged**")
        sb = get_supabase()
        if sb:
            st.success("☁ Supabase connected", icon="✅")
        else:
            st.warning("⚠ No DB — session only", icon="⚠")
        st.markdown("---")
        st.markdown("#### Navigation")
        st.caption("Use the tabs above to navigate between Dashboard, Daily Log, Report, and Settings.")
        st.markdown("---")
        st.markdown("#### Legend")
        for cat in CATS:
            st.markdown(
                f'<span style="color:{cat["hex"]}; font-weight:700;">●</span> '
                f'<span style="color:#e6edf3;">{cat["label"]}</span>',
                unsafe_allow_html=True)
        st.markdown("---")
        tz = get_browser_timezone()
        st.caption(f"Today: {fmt_date(today_str())}  ·  Timezone: {tz}")

# ══════════════════════════════════════════════════════════════════════════════
# TAB 1 — DASHBOARD
# ══════════════════════════════════════════════════════════════════════════════
def render_dashboard(store):
    rigs  = store["rigs"]
    logs  = store["logs"]
    today = today_str()
    t_log = logs.get(today, {})

    # Header
    st.markdown("""
    <div style="background:#161b22;border:1px solid #30363d;border-radius:8px;padding:14px 20px;margin-bottom:16px;">
      <div style="color:#8b949e;font-size:9pt;font-weight:700;letter-spacing:2px;">CONTINUOUS DELIVERY FOR QUALITY AND COMPLIANCE</div>
      <div style="color:#e6edf3;font-size:20pt;font-weight:800;margin-top:2px;">HIL — TEST ENVIRONMENTS</div>
    </div>""", unsafe_allow_html=True)

    # Active / Inactive / In Progress banner
    active_c   = sum(1 for r in rigs if r.get("status","active")=="active")
    inactive_c = sum(1 for r in rigs if r.get("status","active")=="inactive")
    inprog_c   = sum(1 for r in rigs if r.get("status","active")=="in_progress")
    st.markdown(f"""
    <div style="background:#388bfd;border-radius:8px;padding:16px 20px;text-align:center;margin-bottom:16px;">
      <div style="color:rgba(255,255,255,0.9);font-size:13pt;font-weight:600;">Active  /  Inactive  /  In Progress</div>
      <div style="color:#ffffff;font-size:28pt;font-weight:800;letter-spacing:2px;margin-top:4px;">
        {active_c}  /  {inactive_c}  /  {inprog_c}
      </div>
    </div>""", unsafe_allow_html=True)

    # Fleet KPI metrics
    totals = {k:0.0 for k in CAT_KEYS}
    for rig in rigs:
        e = t_log.get(rig["id"],{})
        for k in CAT_KEYS: totals[k] += float(e.get(k,0) or 0)
    active_h  = totals["autoExec"]+totals["manExec"]+totals["autoDev"]
    max_fleet = len(rigs)*DAY_HOURS
    util = round(active_h/max_fleet*100) if max_fleet else 0
    cols = st.columns(6)
    for col, (lbl, val, delta_col) in zip(cols, [
        ("AUTO EXEC",   f"{totals['autoExec']:.1f}h", C_AUTO),
        ("MANUAL EXEC", f"{totals['manExec']:.1f}h",  C_MAN),
        ("AUTO DEV",    f"{totals['autoDev']:.1f}h",  C_DEV),
        ("IDLE",        f"{totals['idle']:.1f}h",     C_IDLE),
        ("DOWN",        f"{totals['down']:.1f}h",     C_DOWN),
        ("FLEET UTIL",  f"{util}%",
         C_GREEN if util>=70 else C_YELLOW if util>=40 else C_RED),
    ]):
        with col:
            st.markdown(f"""
            <div style="background:#161b22;border:1px solid #30363d;border-radius:8px;
                        padding:12px 16px;text-align:center;">
              <div style="color:#8b949e;font-size:8pt;font-weight:700;letter-spacing:1px;">{lbl}</div>
              <div style="color:{delta_col};font-size:22pt;font-weight:800;">{val}</div>
            </div>""", unsafe_allow_html=True)

    st.markdown("---")

    # Weekly chart
    fig = weekly_util_chart(store)
    st.pyplot(fig, use_container_width=True)
    plt.close(fig)

    st.markdown("---")
    st.markdown("#### Rig Status — Today")

    # Rig cards
    n_cols = min(len(rigs), 4)
    rows = [rigs[i:i+n_cols] for i in range(0, len(rigs), n_cols)]
    for row in rows:
        cols = st.columns(len(row))
        for col, rig in zip(cols, row):
            e      = t_log.get(rig["id"], {})
            status = rig.get("status","active")
            note   = rig.get("note","")
            s_col  = STATUS_COLS.get(status, "#888899")
            s_lbl  = status.replace("_"," ").upper()
            active_h_rig = sum(float(e.get(k,0) or 0)
                               for k in ["autoExec","manExec","autoDev"])
            with col:
                hours_html = ""
                for cat in CATS:
                    v=float(e.get(cat["key"],0) or 0)
                    if v>0:
                        hours_html+=f'<div style="color:{cat["hex"]};font-size:9pt;font-weight:700;margin-top:4px;">{cat["short"]} {v:.1f}h</div>'
                note_html = f'<span style="color:#8b949e;font-size:8pt;font-weight:600;line-height:1.3;">• {note}</span>' if note else '<span style="color:transparent;font-size:8pt;line-height:1.3;">placeholder</span>'
                sw_version = e.get("swVersion", e.get("project", ""))
                sw_html = f'<span style="color:#ffffff;font-size:9pt;font-weight:700;">SW: {sw_version}</span>' if sw_version else ""
                st.markdown(f"""
                <div style="background:#161b22;border:1px solid #30363d;
                    border-left:3px solid {s_col};border-radius:8px;
                    padding:14px 16px;margin-bottom:16px;min-height:148px;
                    display:grid;grid-template-rows:auto 1fr auto;gap:12px;">
                  <div style="display:flex;justify-content:space-between;align-items:center;">
                    <span style="color:#e6edf3;font-size:13pt;font-weight:800;">{rig['name']}</span>
                    <span style="color:{s_col};background:rgba(0,0,0,0.3);border:1px solid {s_col}44;
                          border-radius:4px;padding:2px 8px;font-size:8pt;font-weight:700;letter-spacing:0.5px;">{s_lbl}</span>
                  </div>
                  <div style="display:flex;flex-direction:column;justify-content:flex-start;">
                    {hours_html if hours_html else '<div style="color:#484f58;font-size:9pt;">No entry today</div>'}
                  </div>
                  <div style="display:flex;justify-content:space-between;align-items:center;">
                    <div>{note_html}</div>
                    <div>{sw_html}</div>
                  </div>
                </div>""", unsafe_allow_html=True)

    # 7-day history
    st.markdown("---")
    st.markdown("#### 7-Day History")
    dates = last_n_dates(7)
    rows_data = []
    for rig in rigs:
        row = {"RIG": rig["name"]}
        for d in dates:
            e = logs.get(d,{}).get(rig["id"],{})
            if e:
                ah=sum(float(e.get(k,0) or 0) for k in ["autoExec","manExec","autoDev"])
                dh=float(e.get("down",0) or 0)
                row[fmt_short(d)] = f"{ah:.1f}h" if ah>0 else ("DOWN" if dh>0 else "IDLE")
            else:
                row[fmt_short(d)] = "—"
        rows_data.append(row)
    import pandas as pd
    df = pd.DataFrame(rows_data).set_index("RIG")
    st.dataframe(df, use_container_width=True)

    # ── Engineer Handover Board ───────────────────────────────────────────────
    st.markdown("---")
    st.markdown("#### 🔄 Engineer Handover — Today")

    # Collect per-rig engineer assignments
    rig_engineers = []
    for rig in rigs:
        e        = t_log.get(rig["id"], {})
        offshore = e.get("offshore","").strip()
        onsite   = e.get("onsite","").strip()
        if offshore or onsite:
            rig_engineers.append((rig, offshore, onsite))

    # Global handover note
    global_note = t_log.get("_meta",{}).get("handover","").strip()

    if rig_engineers:
        # Column headers
        st.markdown(
            '<div style="display:grid;grid-template-columns:120px 1fr 40px 1fr;'
            'gap:8px;padding:6px 12px;background:#161b22;border-radius:6px;margin-bottom:8px;">'
            '<span style="color:#484f58;font-size:8pt;font-weight:700;letter-spacing:1px;">RIG</span>'
            '<span style="color:#388bfd;font-size:8pt;font-weight:700;letter-spacing:1px;">🌐 OFFSHORE ENGINEER</span>'
            '<span></span>'
            '<span style="color:#3fb950;font-size:8pt;font-weight:700;letter-spacing:1px;">🏢 ONSITE ENGINEER</span>'
            '</div>', unsafe_allow_html=True)

        for rig, offshore, onsite in rig_engineers:
            off_html = (
                f'<div style="background:rgba(56,139,253,0.10);border:1px solid rgba(56,139,253,0.3);'
                f'border-radius:6px;padding:10px 14px;">'
                f'<div style="color:#388bfd;font-size:7.5pt;font-weight:700;letter-spacing:1.5px;">🌐 OFFSHORE</div>'
                f'<div style="color:#e6edf3;font-size:14pt;font-weight:800;">'
                f'{offshore if offshore else "<span style='color:#484f58'>—</span>"}</div>'
                f'</div>')
            on_html = (
                f'<div style="background:rgba(63,185,80,0.10);border:1px solid rgba(63,185,80,0.3);'
                f'border-radius:6px;padding:10px 14px;">'
                f'<div style="color:#3fb950;font-size:7.5pt;font-weight:700;letter-spacing:1.5px;">🏢 ONSITE</div>'
                f'<div style="color:#e6edf3;font-size:14pt;font-weight:800;">'
                f'{onsite if onsite else "<span style='color:#484f58'>—</span>"}</div>'
                f'</div>')
            st.markdown(
                f'<div style="display:grid;grid-template-columns:120px 1fr 40px 1fr;'
                f'gap:8px;align-items:center;margin-bottom:8px;'
                f'background:#161b22;border:1px solid #30363d;border-radius:8px;padding:8px 12px;">'
                f'<span style="color:#e6edf3;font-size:12pt;font-weight:800;">{rig["name"]}</span>'
                f'{off_html}'
                f'<div style="text-align:center;color:#484f58;font-size:18pt;font-weight:300;">→</div>'
                f'{on_html}'
                f'</div>', unsafe_allow_html=True)

        if global_note:
            st.markdown(
                f'<div style="background:rgba(210,153,34,0.08);border:1px solid rgba(210,153,34,0.3);'
                f'border-radius:6px;padding:10px 16px;margin-top:8px;">'
                f'<span style="color:#d29922;font-weight:700;">📋 Handover Note: </span>'
                f'<span style="color:#e6edf3;">{global_note}</span>'
                f'</div>', unsafe_allow_html=True)
    else:
        st.markdown(
            '<div style="color:#484f58;font-size:10pt;font-style:italic;padding:8px 0;">'
            'No engineer assignments for today. Add Offshore / Onsite names per rig in the Daily Log tab.'
            '</div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# TAB 2 — DAILY LOG
# ══════════════════════════════════════════════════════════════════════════════
def render_daily_log(store):
    st.markdown("### ✏ Daily Log Entry")
    col1, col2, col3 = st.columns([2,1,3])
    with col1:
        sel_date = st.date_input("Log Date", value=today_date(),
                                 max_value=today_date(), key="log_date")
    date_str = sel_date.isoformat()
    existing = store["logs"].get(date_str,{})
    with col2:
        st.markdown("<br>", unsafe_allow_html=True)
        if date_str in store["logs"]:
            st.success("✓ Saved")
    with col3:
        st.markdown("<br>", unsafe_allow_html=True)
        st.info(f"Rules: AE ≤{DAY_HOURS}h  ·  ME ≤{HUMAN_POOL}h  ·  AD+ME ≤{HUMAN_POOL}h  ·  Total ≤{DAY_HOURS}h  ·  Idle/Down reason required")

    # Global handover note for the day
    existing_meta = existing.get("_meta", {})
    st.markdown("---")
    hov_col1, hov_col2 = st.columns([3,3])
    global_handover = hov_col1.text_input(
        "📋 Day Handover Note (optional)",
        existing_meta.get("handover",""),
        placeholder="Brief note about today's overall handover…",
        key="global_handover")
    st.markdown("---")
    entries = {}
    has_errors = False

    for rig in store["rigs"]:
        rid = rig["id"]
        e   = existing.get(rid, blank_entry())
        with st.expander(f"**{rig['name']}**  {'✓' if rid in existing else '○'}", expanded=False):
            c1,c2,c3,c4,c5 = st.columns(5)
            ae = c1.number_input("Auto Exec (h)", 0.0, float(DAY_HOURS), float(e.get("autoExec",0)), 0.5, key=f"ae_{rid}")
            me = c2.number_input("Manual Exec (h)",0.0, float(HUMAN_POOL), float(e.get("manExec", 0)), 0.5, key=f"me_{rid}")
            ad = c3.number_input("Auto Dev (h)",  0.0, float(AUTDEV_MAX),  float(e.get("autoDev", 0)), 0.5, key=f"ad_{rid}")
            ih = c4.number_input("Idle (h)",      0.0, float(DAY_HOURS),   float(e.get("idle",    0)), 0.5, key=f"idle_{rid}")
            dh = c5.number_input("Down (h)",      0.0, float(DAY_HOURS),   float(e.get("down",    0)), 0.5, key=f"down_{rid}")

            c6,c7 = st.columns(2)
            sw_version = c6.text_input("SW Version", e.get("swVersion", e.get("project", "")), key=f"sw_{rid}")
            notes = c7.text_input("Notes",   e.get("notes",""),   key=f"notes_{rid}")

            # Engineer handover fields per rig
            st.markdown(
                '<div style="display:flex;align-items:center;gap:8px;margin-top:6px;">'
                '<span style="color:#388bfd;font-size:8pt;font-weight:700;letter-spacing:1px;">🌐 OFFSHORE / ONSITE ENGINEER ASSIGNMENT</span>'
                '</div>', unsafe_allow_html=True)
            ec1, ec2 = st.columns(2)
            offshore = ec1.text_input(
                "🌐 Offshore Engineer", e.get("offshore",""),
                placeholder="Name of offshore engineer…",
                key=f"off_{rid}")
            onsite = ec2.text_input(
                "🏢 Onsite Engineer", e.get("onsite",""),
                placeholder="Name of onsite engineer…",
                key=f"on_{rid}")

            # Reasons
            idle_reason = e.get("idleReason","")
            down_reason = e.get("downReason","")
            if ih > 0:
                r1,r2 = st.columns(2)
                with r1:
                    ir_choice = st.selectbox("Idle Reason ⚠ required",
                        [""] + IDLE_PRESETS, key=f"ir_{rid}",
                        index=(IDLE_PRESETS.index(idle_reason)+1) if idle_reason in IDLE_PRESETS else 0)
                    if ir_choice == "Other" or (idle_reason and idle_reason not in IDLE_PRESETS):
                        ir_custom = st.text_input("Describe idle reason",
                            idle_reason if idle_reason not in IDLE_PRESETS else "",
                            key=f"ir_custom_{rid}")
                        idle_reason = ir_custom
                    else:
                        idle_reason = ir_choice
                    if ih > 0 and not idle_reason:
                        st.error("Idle reason is required")
                        has_errors = True
            if dh > 0:
                r1,r2 = st.columns(2)
                with r1:
                    dr_choice = st.selectbox("Down Reason ⚠ required",
                        [""] + DOWN_PRESETS, key=f"dr_{rid}",
                        index=(DOWN_PRESETS.index(down_reason)+1) if down_reason in DOWN_PRESETS else 0)
                    if dr_choice == "Other" or (down_reason and down_reason not in DOWN_PRESETS):
                        dr_custom = st.text_input("Describe down reason",
                            down_reason if down_reason not in DOWN_PRESETS else "",
                            key=f"dr_custom_{rid}")
                        down_reason = dr_custom
                    else:
                        down_reason = dr_choice
                    if dh > 0 and not down_reason:
                        st.error("Down reason is required")
                        has_errors = True

            # Validation
            warns = []
            if me > HUMAN_POOL:      warns.append(f"ME = {me:.1f}h > {HUMAN_POOL}h max")
            if ad + me > HUMAN_POOL: warns.append(f"AD+ME = {ad+me:.1f}h > {HUMAN_POOL}h pool")
            if ad > AUTDEV_MAX:      warns.append(f"AD = {ad:.1f}h > {AUTDEV_MAX}h max")
            total = ae+me+ad+ih+dh
            if total > DAY_HOURS:  warns.append(f"Total {total:.1f}h > {DAY_HOURS}h")
            if warns:
                st.warning(" · ".join(warns)); has_errors = True
            else:
                st.markdown(f"**Total: {total:.1f}h / {DAY_HOURS}h**")

            entries[rid] = {
                "autoExec":ae,"manExec":me,"autoDev":ad,
                "idle":ih,"idleReason":idle_reason,
                "down":dh,"downReason":down_reason,
                "swVersion":sw_version,"notes":notes,
                "offshore":offshore,"onsite":onsite,
            }

    st.markdown("---")
    if st.button("💾 Save Log", type="primary", disabled=has_errors):
        entries["_meta"] = {"handover": global_handover}
        store["logs"][date_str] = entries
        db_save_log(date_str, entries)
        st.success(f"✓ Log saved for {fmt_date(date_str)}")
        st.rerun()

# ══════════════════════════════════════════════════════════════════════════════
# TAB 3 — REPORT
# ══════════════════════════════════════════════════════════════════════════════
def render_report(store):
    st.markdown("### 📋 Report")

    # Controls
    c1,c2,c3,c4,c5 = st.columns([2,2,2,2,2])
    with c1:
        from_date = st.date_input("From",
            value=today_date()-datetime.timedelta(days=6),
            max_value=today_date(), key="rep_from")
    with c2:
        to_date = st.date_input("To", value=today_date(),
            max_value=today_date(), key="rep_to")
    with c3:
        st.markdown("<br>",unsafe_allow_html=True)
        gen = st.button("Generate Report", type="primary")
    with c4:
        st.markdown("<br>",unsafe_allow_html=True)
        do_pptx = st.button("⬇ Export to PPTX")
    with c5:
        st.markdown("<br>",unsafe_allow_html=True)
        do_csv = st.button("⬇ CSV")

    # Rig selection
    with st.expander("⚙  Rig Filter — select rigs to include in report", expanded=False):
        st.markdown("Toggle rigs on/off for the report:")
        rig_cols = st.columns(min(len(store["rigs"]), 7))
        sel_rigs = {}
        for i,rig in enumerate(store["rigs"]):
            with rig_cols[i % len(rig_cols)]:
                sel = st.toggle(rig["name"], value=True, key=f"rtog_{rig['id']}")
                sel_rigs[rig["id"]] = sel
        sel_ids = [rid for rid,on in sel_rigs.items() if on]
        n_sel = len(sel_ids)
        st.caption(f"{n_sel}/{len(store['rigs'])} rigs selected")

    from_d = from_date.isoformat()
    to_d   = to_date.isoformat()

    # Always compute (Streamlit reruns on every interaction)
    stats, cdata, dates = compute_stats(store, from_d, to_d, sel_ids)

    # KPI strip
    fu = round(sum(r["util"] for r in stats)/len(stats)) if stats else 0
    uc = C_GREEN if fu>=70 else C_YELLOW if fu>=40 else C_RED
    top_rig = max(stats, key=lambda r:r["util"]) if stats else None
    kpi_cols = st.columns(7)
    for col,(lbl,val,c) in zip(kpi_cols,[
        ("PERIOD",       f"{len(dates)}d",                                C_ACCENT),
        ("FLEET UTIL",   f"{fu}%",                                        uc),
        ("AUTO EXEC",    f"{sum(r['autoExec'] for r in stats):.1f}h",    C_AUTO),
        ("MANUAL EXEC",  f"{sum(r['manExec'] for r in stats):.1f}h",     C_MAN),
        ("AUTO DEV",     f"{sum(r['autoDev'] for r in stats):.1f}h",     C_DEV),
        ("DOWN/INACTIVE",f"{sum(r['down'] for r in stats):.1f}h",        C_DOWN),
        ("TOP RIG",      top_rig["name"] if top_rig else "—",            C_ACCENT),
    ]):
        with col:
            st.markdown(f"""
            <div style="background:#161b22;border:1px solid #30363d;border-radius:8px;
                        padding:10px 14px;text-align:center;margin-bottom:8px;">
              <div style="color:#8b949e;font-size:7.5pt;font-weight:700;letter-spacing:1px;">{lbl}</div>
              <div style="color:{c};font-size:18pt;font-weight:800;">{val}</div>
            </div>""", unsafe_allow_html=True)

    # View tabs
    vt = st.tabs(["📊 Summary Table", "🔍 Key Highlights", "📈 Charts", "📅 Daily Detail"])

    # ── Summary Table ─────────────────────────────────────────────────────────
    with vt[0]:
        if not stats:
            st.info("No data for selected period")
        else:
            import pandas as pd
            rows = []
            for r in stats:
                uc2 = "🟢" if r["util"]>=70 else "🟡" if r["util"]>=40 else "🔴"
                rows.append({
                    "RIG": r["name"], "DAYS": r["days"],
                    "AUTO EXEC (h)": f"{r['autoExec']:.1f}",
                    "MANUAL EXEC (h)": f"{r['manExec']:.1f}",
                    "AUTO DEV (h)": f"{r['autoDev']:.1f}",
                    "IDLE (h)": f"{r['idle']:.1f}",
                    "DOWN (h)": f"{r['down']:.1f}",
                    "UTIL": f"{uc2} {r['util']}%",
                })
            # Fleet totals row
            rows.append({
                "RIG":"FLEET TOTAL","DAYS":"—",
                "AUTO EXEC (h)": f"{sum(r['autoExec'] for r in stats):.1f}",
                "MANUAL EXEC (h)": f"{sum(r['manExec'] for r in stats):.1f}",
                "AUTO DEV (h)": f"{sum(r['autoDev'] for r in stats):.1f}",
                "IDLE (h)": f"{sum(r['idle'] for r in stats):.1f}",
                "DOWN (h)": f"{sum(r['down'] for r in stats):.1f}",
                "UTIL": f"{fu}%",
            })
            df=pd.DataFrame(rows).set_index("RIG")
            st.dataframe(df, use_container_width=True)

    # ── Key Highlights ────────────────────────────────────────────────────────
    with vt[1]:
        logs = store["logs"]; rigs_all = store["rigs"]
        idle_g = defaultdict(list); down_g = defaultdict(list)
        for d in dates:
            for rig in [r for r in rigs_all if r["id"] in sel_ids]:
                e=logs.get(d,{}).get(rig["id"],{})
                if not e: continue
                ih=float(e.get("idle",0)or 0); dh=float(e.get("down",0)or 0)
                ir=e.get("idleReason",""); dr=e.get("downReason","")
                if ih>0 and ir: idle_g[ir].append({"rig":rig["name"],"date":fmt_short(d),"hours":ih})
                if dh>0 and dr: down_g[dr].append({"rig":rig["name"],"date":fmt_short(d),"hours":dh})

        def render_group(groups, color, icon, section):
            if not groups: return
            sg=sorted(groups.items(),key=lambda x:sum(i["hours"] for i in x[1]),reverse=True)
            th=sum(sum(i["hours"] for i in v) for _,v in sg)
            st.markdown(f"<span style='color:{color};font-size:13pt;font-weight:700;'>{icon} {section} — {th:.1f}h total</span>",
                        unsafe_allow_html=True)
            for reason,entries in sg:
                total=sum(i["hours"] for i in entries)
                affected=list(set(i["rig"] for i in entries))
                chips=" &nbsp; ".join(f'<span style="background:#21262d;border:1px solid #30363d;border-radius:4px;padding:2px 8px;font-size:9pt;color:#c9d1d9;">{i["rig"]} · {i["date"]} · {i["hours"]:.1f}h</span>' for i in entries)
                st.markdown(f"""
                <div style="background:#161b22;border:1px solid {color}33;border-left:3px solid {color};
                    border-radius:8px;padding:12px 16px;margin:8px 0;">
                  <div style="display:flex;justify-content:space-between;">
                    <span style="color:{color};font-weight:700;font-size:11pt;">{icon} {reason}</span>
                    <span style="color:#8b949e;font-size:9pt;">{total:.1f}h · {len(entries)} occurrence(s)</span>
                  </div>
                  <div style="margin-top:8px;">{chips}</div>
                  <div style="color:#8b949e;font-size:9pt;margin-top:6px;">Rigs: {', '.join(affected)}</div>
                </div>""", unsafe_allow_html=True)

        render_group(down_g, C_RED, "⚠", "DOWN / INACTIVE EVENTS")
        st.markdown("") 
        render_group(idle_g, C_YELLOW, "⏸", "IDLE EVENTS")
        if not idle_g and not down_g:
            st.info("No idle or down events in this period.")

    # ── Charts ────────────────────────────────────────────────────────────────
    with vt[2]:
        if not stats:
            st.info("No data — click Generate Report")
        else:
            fig1, fig2 = report_charts(stats, cdata)
            st.pyplot(fig1, use_container_width=True)
            st.pyplot(fig2, use_container_width=True)
            plt.close(fig1); plt.close(fig2)

    # ── Daily Detail ──────────────────────────────────────────────────────────
    with vt[3]:
        if not dates:
            st.info("No dates in range")
        else:
            import pandas as pd
            logs = store["logs"]
            rigs_sel = [r for r in store["rigs"] if r["id"] in sel_ids]
            rows=[]
            for d in dates:
                row={"DATE":fmt_short(d)}
                for rig in rigs_sel:
                    e=logs.get(d,{}).get(rig["id"],{})
                    for cat in CATS:
                        v=float(e.get(cat["key"],0)or 0)
                        row[f"{rig['name']}\n{cat['short']}"]=f"{v:.1f}" if v>0 else "·"
                rows.append(row)
            df=pd.DataFrame(rows).set_index("DATE")
            st.dataframe(df, use_container_width=True)

    # ── PPTX / CSV exports ───────────────────────────────────────────────────
    if do_pptx and stats:
        fig1, fig2 = report_charts(stats, cdata)
        buf = build_pptx(store, stats, cdata, dates, from_d, to_d, fig1, fig2)
        plt.close("all")
        st.download_button("📥 Download PPTX", buf,
            file_name=f"HIL_Report_{today_str()}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")

    if do_csv and stats:
        import csv, io as _io
        buf = _io.StringIO()
        w = csv.writer(buf)
        w.writerow(["RIG","DAYS","AUTO_EXEC_H","MANUAL_EXEC_H","AUTO_DEV_H","IDLE_H","DOWN_H","UTIL_PCT"])
        for r in stats:
            w.writerow([r["name"],r["days"],f"{r['autoExec']:.1f}",f"{r['manExec']:.1f}",
                        f"{r['autoDev']:.1f}",f"{r['idle']:.1f}",f"{r['down']:.1f}",f"{r['util']}"])
        st.download_button("📥 Download CSV", buf.getvalue().encode(),
            file_name=f"HIL_Report_{today_str()}.csv", mime="text/csv")

# ══════════════════════════════════════════════════════════════════════════════
# TAB 4 — SETTINGS
# ══════════════════════════════════════════════════════════════════════════════
def render_settings(store):
    st.markdown("### ⚙ Rig Management")
    c1, c2 = st.columns([1,1])

    with c1:
        st.markdown("#### Current Rigs")
        for i,rig in enumerate(store["rigs"]):
            with st.expander(f"**{rig['name']}** — {rig.get('status','active').replace('_',' ').upper()}", expanded=False):
                new_name = st.text_input("Name", rig["name"], key=f"sname_{rig['id']}")
                new_status = st.selectbox("Status", STATUS_OPTS,
                    index=STATUS_OPTS.index(rig.get("status","active")),
                    key=f"sstatus_{rig['id']}")
                new_note = st.text_input("Note / Issue", rig.get("note",""), key=f"snote_{rig['id']}")
                sc1, sc2 = st.columns(2)
                if sc1.button("Update", key=f"supdate_{rig['id']}"):
                    store["rigs"][i]["name"]   = new_name
                    store["rigs"][i]["status"] = new_status
                    store["rigs"][i]["note"]   = new_note
                    db_save_rigs(store["rigs"])
                    st.success("Updated ✓"); st.rerun()
                if sc2.button("Remove", key=f"sremove_{rig['id']}",
                              help="Historical data is kept"):
                    store["rigs"].pop(i)
                    db_save_rigs(store["rigs"])
                    st.success("Removed ✓"); st.rerun()

        st.markdown("---")
        st.markdown("#### Add New Rig")
        new_nm = st.text_input("Rig name", placeholder=f"HIL-{len(store['rigs'])+1:02d}",
                               key="snew_nm")
        new_st = st.selectbox("Status", STATUS_OPTS, key="snew_status")
        new_nt = st.text_input("Note (optional)", key="snew_note")
        if st.button("+ Add Rig", type="primary"):
            nm = new_nm.strip() or f"HIL-{len(store['rigs'])+1:02d}"
            store["rigs"].append({"id":f"rig-{int(datetime.datetime.now().timestamp())}",
                                   "name":nm,"status":new_st,"note":new_nt})
            db_save_rigs(store["rigs"])
            st.success(f"Added {nm} ✓"); st.rerun()



    with c2:
        st.markdown("#### Hour Rules — 24h Day")
        for cat in CATS:
            label = " — <= 24h" if cat["key"] == "autoExec" else " — <= 18h" if cat["key"] in ["manExec","autoDev"] else " — <= 24h, reason required"
            st.markdown(
                f'<span style="color:{cat["hex"]};font-weight:700;font-size:11pt;">● {cat["label"]}</span>'
                + label,
                unsafe_allow_html=True)
        st.markdown(f"""
        <div style="background:#161b22;border:1px solid #30363d;border-radius:6px;padding:12px 16px;margin-top:12px;">
          <div style="color:#388bfd;font-weight:700;">AD + ME share {HUMAN_POOL}h human pool</div>
          <div style="color:#8b949e;font-size:9pt;margin-top:4px;">
            Auto Dev + Manual Exec combined ≤ {HUMAN_POOL}h per rig per day
          </div>
        </div>""", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════════
def main():
    init_state()
    store = st.session_state.store

    render_sidebar(store)

    # App header
    now = get_user_now()
    st.markdown(f"""
    <div style="display:flex;justify-content:space-between;align-items:center;
                border-bottom:1px solid #30363d;padding-bottom:12px;margin-bottom:0;">
      <div>
        <div style="color:#8b949e;font-size:9pt;font-weight:700;letter-spacing:2px;">HIL RIG  ·  UTILISATION TRACKER</div>
        <div style="color:#e6edf3;font-size:20pt;font-weight:800;margin-top:2px;">RIG STATUS & USAGE BOARD</div>
      </div>
      <div style="text-align:right;color:#8b949e;font-size:9pt;">
        {now.strftime("%a, %d %b %Y %H:%M %Z")}<br>
        {len(store['rigs'])} rigs · {len(store['logs'])} days logged
      </div>
    </div>""", unsafe_allow_html=True)

    tab1, tab2, tab3, tab4 = st.tabs([
        "  📊  Dashboard  ",
        "  ✏  Daily Log  ",
        "  📋  Report  ",
        "  ⚙  Settings  ",
    ])

    with tab1: render_dashboard(store)
    with tab2: render_daily_log(store)
    with tab3: render_report(store)
    with tab4: render_settings(store)

if __name__ == "__main__":
    main()
